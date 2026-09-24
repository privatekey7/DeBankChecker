"""
Многопроходная сверка балансов: N полных независимых проходов по всем
кошелькам и итоговый отчёт стабильности.

    python audit_passes.py [--passes 5] [--workers 0] [--direct] [--max-extra 3]

Каждый проход = check_wallet для каждого адреса (внутри — минимум
CORROBORATION_MIN_AGREE согласованных выборок). Кошелёк СТАБИЛЕН, если у него
≥ N результатов OK и все они сошлись в пределах допуска. Для остальных
выполняются дополнительные точечные проходы (до --max-extra).

Результаты: Results/audit_<ts>/
    pass_<k>.jsonl      — результаты каждого прохода (без tokens/protocols/nft)
    audit_report.xlsx   — сводка по кошелькам: значения по проходам, разброс, вердикт
    audit_report.csv    — то же в CSV
    final_results.json  — итоговый результат по каждому кошельку (последний OK)
    debank_*.xlsx       — стандартный Excel-экспорт по финальным результатам

Код возврата: 0 — все кошельки стабильны, 2 — есть нестабильные.
--direct: без прокси (по умолчанию прокси обязательны, как и в main.py).
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--passes", type=int, default=5)
    ap.add_argument("--workers", type=int, default=0, help="0 = как main.py: min(500, кошельки, прокси×5); без прокси 8")
    ap.add_argument("--direct", action="store_true", help="работать без прокси")
    ap.add_argument("--max-extra", type=int, default=3, help="доп. проходов для нестабильных кошельков")
    ap.add_argument("--wallets", default=None)
    ap.add_argument("--limit", type=int, default=0, help="взять первые N кошельков (отладка)")
    args = ap.parse_args()

    if args.direct:
        os.environ["DBC_DIRECT"] = "1"

    # импорт после установки окружения: config читает DBC_DIRECT при загрузке
    from debank_checker import __version__, config
    from debank_checker.checker import _agree, wait_nfts
    from debank_checker.export.excel import export_to_excel
    from debank_checker.proxy.manager import ProxyManager, drop_dead_proxies, load_proxies
    from main import check_all, load_wallets, setup_encoding, worker_count

    setup_encoding()
    wallets = load_wallets(args.wallets or config.WALLETS_FILE)
    if args.limit:
        wallets = wallets[:args.limit]
    proxies = [] if args.direct else load_proxies(config.PROXIES_FILE)
    if not wallets:
        print("Нет кошельков")
        return 1
    if not proxies and not config.ALLOW_DIRECT:
        print("Прокси обязательны (или --direct)")
        return 1
    if proxies:
        total_proxies = len(proxies)
        proxies, _ = drop_dead_proxies(proxies)  # мёртвые исключаются молча
        if not proxies:
            print("Ни один прокси не работает")
            return 1
    workers = args.workers if args.workers > 0 else worker_count(len(wallets), len(proxies))

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = Path(config.OUTPUT_DIR) / f"audit_{ts}"
    out.mkdir(parents=True, exist_ok=True)
    print(f"v{__version__} | кошельков: {len(wallets)} | {f'рабочие прокси: {len(proxies)} из {total_proxies}' if proxies else 'прокси: direct'} | "
          f"проходов: {args.passes} | воркеров: {workers}")
    print(f"Каталог: {out}")

    pm = ProxyManager(proxies)
    results: dict[str, list[dict]] = {w: [] for w in wallets}
    timings: list[tuple[int, int, float, float]] = []  # (проход, кошельков, баланс с, всего с)

    def slim(r: dict) -> dict:
        return {k: v for k, v in r.items() if k not in ("tokens_data", "protocols_data", "nft_data")}

    def run_pass(k: int, targets: list[str]) -> None:
        t0 = time.perf_counter()

        def progress(done: int, total: int) -> None:
            if done % 50 == 0 or done == total:
                print(f"  проход {k}: {done}/{total}  {time.perf_counter() - t0:.1f}s", flush=True)

        rows = check_all(targets, pm, workers, on_done=progress)
        t_balance = time.perf_counter() - t0
        left = wait_nfts(timeout=180)
        t_total = time.perf_counter() - t0
        with open(out / f"pass_{k}.jsonl", "a", encoding="utf-8") as f:
            for addr, r in zip(targets, rows):
                r["pass"] = k
                results[addr].append(r)
                f.write(json.dumps(slim(r), ensure_ascii=False) + "\n")
        timings.append((k, len(targets), t_balance, t_total))
        ok = [r for r in rows if r["status"] == "OK"]
        unv = sum(1 for r in rows if r["status"] == "UNVERIFIED")
        err = sum(1 for r in rows if r["status"] == "ERROR")
        print(f"Проход {k}: OK={len(ok)} UNVERIFIED={unv} ERROR={err} | сумма OK=${sum(r['total_usd'] for r in ok):,.2f}"
              f" | баланс {t_balance:.1f}s, с NFT {t_total:.1f}s"
              + (f" | NFT не догружены у {left}" if left else ""), flush=True)

    def stable(addr: str) -> tuple[bool, list[float]]:
        oks = [r["total_usd"] for r in results[addr] if r["status"] == "OK"]
        if len(oks) < args.passes:
            return False, oks
        base = min(oks)
        return all(_agree(base, v) for v in oks), oks

    for k in range(1, args.passes + 1):
        run_pass(k, wallets)

    extra = 0
    while extra < args.max_extra:
        unstable = [w for w in wallets if not stable(w)[0]]
        if not unstable:
            break
        extra += 1
        print(f"\nДоп. проход {extra}: нестабильных/недобранных кошельков {len(unstable)}")
        run_pass(args.passes + extra, unstable)

    # --- отчёт ---
    max_n = max(len(v) for v in results.values())
    rows = []
    final: list[dict] = []
    n_stable = 0
    for w in wallets:
        ok_vals = [r["total_usd"] for r in results[w] if r["status"] == "OK"]
        is_stable, _ = stable(w)
        n_stable += is_stable
        last_ok = next((r for r in reversed(results[w]) if r["status"] == "OK"), None)
        chosen = last_ok or results[w][-1]
        if not is_stable and chosen["status"] == "OK":
            chosen = dict(chosen)
            chosen["status"] = "UNVERIFIED"
            chosen["error"] = (f"не набрано {args.passes} согласованных OK-проходов: "
                               + ", ".join(f"${v:.2f}" for v in ok_vals)) + ("; " + chosen["error"] if chosen.get("error") else "")
        final.append(chosen)
        row = {"address": w, "verdict": "STABLE" if is_stable else "UNSTABLE",
               "ok_passes": len(ok_vals), "total_passes": len(results[w]),
               "min": round(min(ok_vals), 2) if ok_vals else None,
               "max": round(max(ok_vals), 2) if ok_vals else None,
               "spread": round(max(ok_vals) - min(ok_vals), 2) if ok_vals else None,
               "tokens_usd": round(chosen.get("tokens_usd", 0), 2),
               "protocols_usd": round(chosen.get("protocols_usd", 0), 2),
               "nft_collections": len(chosen.get("nft_data") or []),
               "aggregate_rejected": sum(1 for r in results[w] if not r.get("aggregate_agrees", True)),
               "rabby_aggregate_last": (round(chosen["aggregate_usd"], 2) if chosen.get("aggregate_usd") is not None else None),
               "unverified_usd": round(chosen.get("unverified_usd", 0), 2),
               "note": chosen.get("error", "") if chosen["status"] != "OK" else "; ".join(chosen.get("notes", []))}
        for i in range(max_n):
            r = results[w][i] if i < len(results[w]) else None
            row[f"pass_{i + 1}"] = (round(r["total_usd"], 2) if r and r["status"] == "OK"
                                    else (r["status"] if r else ""))
        rows.append(row)

    fieldnames = list(rows[0].keys())
    with open(out / "audit_report.csv", "w", encoding="utf-8", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=fieldnames)
        wr.writeheader()
        wr.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Audit"
        ws.append(fieldnames)
        for c in ws[1]:
            c.font = Font(bold=True)
        red = PatternFill("solid", fgColor="FCE8E6")
        green = PatternFill("solid", fgColor="E6F4EA")
        for row in rows:
            ws.append([row[k] for k in fieldnames])
            for c in ws[ws.max_row]:
                c.fill = green if row["verdict"] == "STABLE" else red
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions[get_column_letter(fieldnames.index("note") + 1)].width = 80
        wb.save(out / "audit_report.xlsx")
    except Exception as e:  # noqa: BLE001 — openpyxl не обязателен для отчёта
        print(f"xlsx-отчёт не записан: {e}")

    with open(out / "final_results.json", "w", encoding="utf-8") as f:
        json.dump([slim(r) for r in final], f, ensure_ascii=False, indent=1)
    xlsx = export_to_excel(final, out)

    final_ok = [r for r in final if r["status"] == "OK"]
    print("\n================ ИТОГ ================")
    for k, n, t_bal, t_all in timings:
        vals = [next((r["total_usd"] for r in results[w] if r.get("pass") == k and r["status"] == "OK"), None)
                for w in wallets]
        ok_n = sum(1 for v in vals if v is not None)
        print(f"  проход {k}: OK {ok_n}/{n}  сумма ${sum(v for v in vals if v is not None):,.2f}"
              f"  |  баланс {t_bal:.1f}s, с NFT {t_all:.1f}s")
    print(f"  стабильных кошельков ({args.passes}+ согласованных OK): {n_stable}/{len(wallets)}")
    print(f"  итоговая сумма (OK): ${sum(r['total_usd'] for r in final_ok):,.2f}  ({len(final_ok)} кошельков)")
    print(f"  отчёт: {out / 'audit_report.xlsx'}\n  Excel: {xlsx}")
    return 0 if n_stable == len(wallets) else 2


if __name__ == "__main__":
    sys.exit(main())
